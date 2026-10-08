#!/usr/bin/env python3
"""
Проверка книги сметы Ава Тетис перед КП · канон/инструменты/parse_smeta.py · 27.09.2026

Что делает:
  1) лист «Для КП»: этапы и строки работ; в КП идут строки со стоимостью больше нуля;
  2) лист расчёта, на который ссылается «Для КП» (обычно «Расчёты (с разработкой)»);
  3) календарь платежей с этого листа («Прогнозный календарь платежей …»),
     а если его там нет — таблица «№ п/п · Наименование · Сумма …» того же листа;
  4) проверки: суммы этапов и ИТОГО, строки работ, график платежей по канону (блок 5),
     названия этапов по канону (блок 4).

Запуск:
  python3 parse_smeta.py СМЕТА.xlsx                  # сводка в консоль
  python3 parse_smeta.py СМЕТА.xlsx --md сводка.md    # и в файл
  python3 parse_smeta.py СМЕТА.xlsx --json сводка.json

Код выхода: 0 — проверки прошли; 1 — есть что править в смете; 2 — книгу не удалось разобрать.

Читает сохранённые значения формул. Если книга собрана скриптом и ни разу не открывалась
в Excel, значений в ней нет — скрипт скажет об этом: книгу открыть в Excel и сохранить.
Ставку часа и трудозатраты не выводит: в КП их нет (канон).
"""
import argparse
import json
import re
import sys
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook

# Канон, блок 4 «Смета», названия этапов (27.09.2026, как на листах КП-2026-0287).
CANON = {
    1: "Моделирование",
    2: "Настройка штатного функционала",
    3: "Разработка и интеграции",
    4: "Введение в опытно-промышленную эксплуатацию",
}
CODE = re.compile(r"^\d+(?:\.\d+)*\.?$")
EPS = 0.01


def txt(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = txt(v).replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def money(x):
    if x is None:
        return "—"
    s = f"{x:,.2f}".replace(",", " ").replace(".", ",")
    return s.replace("-", "−")


def days_txt(x):
    if x is None:
        return "—"
    s = (f"{x:.0f}" if float(x).is_integer() else f"{x:g}").replace(".", ",")
    return s.replace("-", "−")


def kp_sheet(wb):
    for ws in wb.worksheets:
        if "кп" in ws.title.lower():
            return ws
    raise SystemExit("Не найден лист «Для КП»")


def header(ws):
    """Строка шапки листа «Для КП» и номера колонок по названиям."""
    for row in ws.iter_rows(min_row=1, max_row=20):
        vals = [txt(c.value) for c in row]
        if "№" in vals and "Стоимость" in vals:
            cols = {}
            for i, v in enumerate(vals):
                if v and v not in cols:
                    cols[v] = i + 1
            return row[0].row, cols
    raise SystemExit("На листе «Для КП» не найдена шапка «№ … Стоимость»")


def read_kp(ws):
    hr, cols = header(ws)
    c_code = cols["№"]
    c_name = cols.get("Этап/Задача", c_code + 1)
    c_desc = cols.get("Описание задачи")
    c_qty = cols.get("Кол-во")
    c_cost = cols["Стоимость"]
    items, total, total_row = [], None, None
    for r in range(hr + 1, ws.max_row + 1):
        code = txt(ws.cell(r, c_code).value)
        if code.upper() == "ИТОГО" or txt(ws.cell(r, c_name).value).upper() == "ИТОГО":
            total, total_row = num(ws.cell(r, c_cost).value), r
            break
        if not CODE.match(code):
            continue
        items.append({
            "row": r,
            "code": code.rstrip("."),
            "name": txt(ws.cell(r, c_name).value),
            "desc": txt(ws.cell(r, c_desc).value) if c_desc else "",
            "qty": txt(ws.cell(r, c_qty).value) if c_qty else "",
            "cost": num(ws.cell(r, c_cost).value),
        })
    return items, total, total_row, c_cost


def calc_sheet(wb_f, wb_v, kp_title, c_cost, total):
    """Лист расчёта: тот, на который ссылаются формулы колонки «Стоимость» листа «Для КП»."""
    refs = Counter()
    ws_f = wb_f[kp_title]
    for r in range(1, ws_f.max_row + 1):
        v = ws_f.cell(r, c_cost).value
        if isinstance(v, str) and v.startswith("="):
            for m in re.finditer(r"'([^']+)'!|([A-Za-zА-Яа-яЁё0-9_.]+)!", v):
                refs[(m.group(1) or m.group(2)).strip()] += 1
    for name, _ in refs.most_common():
        if name in wb_v.sheetnames:
            return wb_v[name], "по ссылкам листа «Для КП»"
    for ws in wb_v.worksheets:  # запасной путь — по совпадению ИТОГО
        if not ws.title.lower().startswith("расч"):
            continue
        for row in ws.iter_rows():
            if txt(row[0].value).upper() == "ИТОГО":
                if any(num(c.value) is not None and total is not None and abs(num(c.value) - total) < EPS for c in row):
                    return ws, "по совпадению ИТОГО"
                break
    return None, ""


def find_label(ws, prefix):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip().startswith(prefix):
                return c
    return None


def first_number_right(ws, cell):
    for col in range(cell.column + 1, min(ws.max_column, cell.column + 12) + 1):
        v = num(ws.cell(cell.row, col).value)
        if v is not None:
            return v
    return None


def read_calendar(ws):
    """Календарь «Прогнозный календарь платежей …»: платежи по колонкам «N платеж»."""
    head = find_label(ws, "Прогнозный календарь платежей")
    if head is None:
        return None
    pay_cols = []
    for col in range(head.column + 1, ws.max_column + 1):
        m = re.match(r"^\s*(\d+)\s*платеж", txt(ws.cell(head.row, col).value), re.I)
        if m:
            pay_cols.append((int(m.group(1)), col))
        elif pay_cols:
            break
    rows = {}
    for r in range(head.row + 1, head.row + 8):
        lab = txt(ws.cell(r, head.column).value).lower()
        if lab.startswith("количество дней"):
            rows["days"] = r
        elif lab.startswith("предварительные даты") or lab.startswith("дата"):
            rows["date"] = r
        elif lab.startswith("сумма платежа"):
            rows["sum"] = r
    pays = []
    for n, col in pay_cols:
        d = ws.cell(rows["date"], col).value if "date" in rows else None
        pays.append({
            "n": n,
            "sum": num(ws.cell(rows["sum"], col).value) if "sum" in rows else None,
            "days": num(ws.cell(rows["days"], col).value) if "days" in rows else None,
            "date": txt(d),
        })
    return {"kind": "календарь", "title": txt(head.value), "payments": pays}


def read_table(ws):
    """Таблица «№ п/п · Наименование · Сумма …» (старый план-график, листы ППИ)."""
    for row in ws.iter_rows():
        vals = [txt(c.value) for c in row]
        if "№ п/п" not in vals:
            continue
        i = vals.index("№ п/п")
        hdr = []
        j = i
        while j < len(vals) and vals[j]:
            hdr.append(vals[j])
            j += 1
        if not any(h.startswith("Наименование") for h in hdr):
            continue
        k_sum = next((k for k, h in enumerate(hdr) if h.startswith("Сумма")), None)
        k_date = next((k for k, h in enumerate(hdr) if "дата" in h.lower()), None)
        pays, r = [], row[0].row + 1
        while r <= ws.max_row:
            cells = [ws.cell(r, row[i].column + k).value for k in range(len(hdr))]
            if not any(txt(v) for v in cells):
                break
            pays.append({
                "n": len(pays) + 1,
                "name": txt(cells[1]) if len(cells) > 1 else "",
                "sum": num(cells[k_sum]) if k_sum is not None else None,
                "days": None,
                "date": txt(cells[k_date]) if k_date is not None else "",
            })
            r += 1
        return {"kind": "таблица", "title": "План-график платежей", "payments": pays}
    return None


def analyse(path):
    try:
        wb_v = load_workbook(path, data_only=True)
        wb_f = load_workbook(path, data_only=False)
    except Exception as e:  # noqa: BLE001
        raise SystemExit(f"Книгу не удалось открыть: {e}")
    kp = kp_sheet(wb_v)
    items, total, total_row, c_cost = read_kp(kp)
    if total is None:
        raise SystemExit("На листе «Для КП» нет значения ИТОГО. Если книга собрана скриптом, "
                         "откройте её в Excel и сохраните — тогда появятся значения формул.")

    codes = [it["code"] for it in items]
    stages, rows = {}, []
    for it in items:
        parts = it["code"].split(".")
        st = int(parts[0])
        if len(parts) == 1:
            stages[st] = {"n": st, "name": it["name"], "cost": it["cost"], "rows": [], "canon": CANON.get(st)}
            continue
        it["stage"] = st
        it["container"] = any(c.startswith(it["code"] + ".") for c in codes)
        if not it["container"] and (it["cost"] or 0) > EPS:
            rows.append(it)
            stages.setdefault(st, {"n": st, "name": "", "cost": None, "rows": [], "canon": CANON.get(st)})["rows"].append(it)

    checks = []  # (ok: True/False/None-инфо, текст)

    ssum = sum(s["cost"] or 0 for s in stages.values())
    if abs(ssum - total) < EPS:
        checks.append((True, f"Сумма этапов = ИТОГО = {money(total)}"))
    else:
        checks.append((False, f"Сумма этапов {money(ssum)} ≠ ИТОГО {money(total)}"))

    bad = []
    for s in stages.values():
        rsum = sum(r["cost"] for r in s["rows"])
        if s["cost"] is not None and abs(rsum - s["cost"]) > EPS:
            bad.append(f"этап {s['n']:02d}: {money(s['cost'])} против {money(rsum)} по строкам")
    for it in items:
        if it.get("container") and it["cost"] is not None:
            lsum = sum(x["cost"] or 0 for x in items
                       if x["code"].startswith(it["code"] + ".")
                       and not any(c.startswith(x["code"] + ".") for c in codes))
            if abs(lsum - it["cost"]) > EPS:
                bad.append(f"подраздел {it['code']}: {money(it['cost'])} против {money(lsum)} по задачам")
    if bad:
        checks.append((False, "Суммы не сходятся: " + "; ".join(bad)))
    else:
        checks.append((True, "Каждый этап и подраздел равен сумме своих строк"))

    checks.append((None, f"Строк работ со стоимостью больше нуля: {len(rows)} — это число для графы "
                         f"«Объём» на титуле (подразделы не считаются, чтобы не было двойного счёта)"))

    for n in sorted(stages):
        s = stages[n]
        if n not in CANON:
            checks.append((False, f"Этап {n:02d} «{s['name']}» — в каноне четыре этапа"))
        elif s["name"] and s["name"] != CANON[n]:
            checks.append((None, f"Этап {n:02d} в книге «{s['name']}» → в КП «{CANON[n]}»"))
        if not s["rows"]:
            checks.append((None, f"Этап {n:02d} без работ со стоимостью больше нуля"))

    calc, how = calc_sheet(wb_f, wb_v, kp.title, c_cost, total)
    cal, term = None, None
    if calc is None:
        checks.append((False, "Не найден лист расчёта, на который ссылается «Для КП»"))
    else:
        cal = read_calendar(calc) or read_table(calc)
        lab = find_label(calc, "Срок реализации в днях")
        term = first_number_right(calc, lab) if lab else None
    if calc is not None and cal is None:
        checks.append((False, f"На листе «{calc.title}» не найден календарь платежей"))
    if cal:
        pays = cal["payments"]
        psum = sum(p["sum"] or 0 for p in pays)
        if abs(psum - total) < EPS:
            checks.append((True, f"Сумма платежей = ИТОГО = {money(total)}"))
        else:
            checks.append((False, f"Сумма платежей {money(psum)} ≠ ИТОГО {money(total)}"))
        live = [p for p in pays if (p["sum"] or 0) > EPS]
        for p in pays:
            if p["sum"] is None:
                checks.append((False, f"Платёж {p['n']}: нет суммы"))
            elif p["sum"] < -EPS:
                checks.append((False, f"Платёж {p['n']}: отрицательная сумма {money(p['sum'])}"))
            elif abs(p["sum"]) <= EPS:
                checks.append((False, f"Платёж {p['n']} нулевой — по канону (блок 5) нулевая строка "
                                      f"сворачивается, платежей станет {len(live)}. Править смету, не КП"))
            if p["days"] is not None and p["n"] > 1:
                if p["days"] < 0:
                    checks.append((False, f"Платёж {p['n']}: отрицательный интервал {days_txt(p['days'])} дн."))
                elif not float(p["days"]).is_integer():
                    checks.append((False, f"Платёж {p['n']}: дробный интервал {days_txt(p['days'])} дн."))
            if (p["sum"] or 0) > EPS and not p["date"].strip():
                checks.append((False, f"Платёж {p['n']}: нет даты"))

    return {
        "file": path,
        "kp_sheet": kp.title,
        "calc_sheet": calc.title if calc is not None else None,
        "calc_sheet_found": how,
        "total": total,
        "term_days": term,
        "stages": [{"n": s["n"], "book": s["name"], "canon": s["canon"], "cost": s["cost"],
                    "rows": len(s["rows"])} for s in (stages[k] for k in sorted(stages))],
        "rows": [{"code": r["code"], "name": r["name"], "qty": r["qty"], "cost": r["cost"]} for r in rows],
        "calendar": cal,
        "checks": [{"ok": ok, "text": t} for ok, t in checks],
    }


def to_md(d):
    out = []
    name = d["file"].replace("\\", "/").split("/")[-1]
    out.append(f"# Смета · {name}\n")
    out.append(f"Лист КП: «{d['kp_sheet']}» · лист расчёта: «{d['calc_sheet'] or '—'}»"
               + (f" ({d['calc_sheet_found']})" if d["calc_sheet_found"] else ""))
    term = f" · срок реализации по книге: {days_txt(d['term_days'])} дн." if d["term_days"] is not None else ""
    out.append(f"ИТОГО работ: {money(d['total'])}{term}\n")
    out.append("## Этапы\n")
    out.append("| Этап | В книге | В КП | Сумма | Строк работ |")
    out.append("|---|---|---|---:|---:|")
    for s in d["stages"]:
        out.append(f"| {s['n']:02d} | {s['book']} | {s['canon'] or '—'} | {money(s['cost'])} | {s['rows']} |")
    out.append(f"\nСтрок работ со стоимостью больше нуля: **{len(d['rows'])}**\n")
    out.append("## Строки работ\n")
    out.append("| Код | Работа | Кол-во | Сумма |")
    out.append("|---|---|---:|---:|")
    for r in d["rows"]:
        out.append(f"| {r['code']} | {r['name']} | {r['qty']} | {money(r['cost'])} |")
    cal = d["calendar"]
    if cal:
        out.append(f"\n## Платежи · {cal['title']}\n")
        out.append("| Платёж | Сумма | Дней с предыдущего | Дата |")
        out.append("|---:|---:|---:|---|")
        for p in cal["payments"]:
            out.append(f"| {p['n']} | {money(p['sum'])} | {days_txt(p['days'])} | {p['date'] or '—'} |")
    out.append("\n## Проверки\n")
    bad = 0
    for c in d["checks"]:
        mark = "✔" if c["ok"] is True else ("✖" if c["ok"] is False else "·")
        bad += c["ok"] is False
        out.append(f"- {mark} {c['text']}")
    out.append("\n**Итог: " + ("проверки прошли.**" if not bad else f"есть что править в смете — {bad}.**"))
    return "\n".join(out) + "\n"


def main():
    ap = argparse.ArgumentParser(description="Проверка книги сметы перед КП")
    ap.add_argument("xlsx")
    ap.add_argument("--md", help="куда сохранить сводку в markdown")
    ap.add_argument("--json", help="куда сохранить данные в JSON")
    a = ap.parse_args()
    try:
        d = analyse(a.xlsx)
    except SystemExit as e:
        print(f"Ошибка: {e}", file=sys.stderr)
        sys.exit(2)
    md = to_md(d)
    sys.stdout.write(md)
    if a.md:
        with open(a.md, "w", encoding="utf-8") as f:
            f.write(md)
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=2)
    sys.exit(1 if any(c["ok"] is False for c in d["checks"]) else 0)


if __name__ == "__main__":
    main()
